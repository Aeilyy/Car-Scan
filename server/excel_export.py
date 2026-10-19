import datetime as dt
import io
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image

COLUMNS = [
    ("Date & time", 18, "created_at"),
    ("Owner", 28, "owner_name"),
    ("VIN (serial number)", 23, "vin"),
    ("Licence plate", 14, "license_plate"),
    ("Description", 42, "description"),
    ("Photo", 36, None),
]
THUMB_MAX_W = 250
THUMB_MAX_H = 190

FONT = Font(name="Arial", size=11)
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="17212B")
TOP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
THIN = Side(style="thin", color="C5CCD3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _thumbnail(path: Path) -> XLImage:
    img = Image.open(path)
    img.thumbnail((THUMB_MAX_W, THUMB_MAX_H))
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=85)
    buf.seek(0)
    xl = XLImage(buf)
    xl.width, xl.height = img.size
    return xl


def build_workbook(rows: List[dict], photo_dir: Path) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vehicles"

    for col, (header, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    r = 2
    for i, entry in enumerate(rows):
        if i > 0:
            r += 1

        for col, (_, _, key) in enumerate(COLUMNS, start=1):
            if key is None:
                continue
            value = entry.get(key, "")
            if key == "created_at":
                value = dt.datetime.fromisoformat(value)
            cell = ws.cell(row=r, column=col, value=value)
            cell.font, cell.alignment, cell.border = FONT, TOP_LEFT, BOX
            if key == "created_at":
                cell.number_format = "dd.mm.yyyy hh:mm"

        photo_cell = ws.cell(row=r, column=len(COLUMNS))
        photo_cell.border = BOX
        photo_path = photo_dir / entry["photo"]
        row_height_pt = 30.0
        if photo_path.exists():
            img = _thumbnail(photo_path)
            ws.add_image(img, photo_cell.coordinate)
            row_height_pt = img.height * 0.75 + 6
        else:
            photo_cell.value = f"Photo missing: {entry['photo']}"
            photo_cell.font = FONT
        ws.row_dimensions[r].height = row_height_pt
        r += 1

    if not rows:
        ws.cell(row=2, column=1, value="No entries yet.").font = FONT

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
